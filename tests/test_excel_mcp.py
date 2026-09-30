import asyncio

import pytest
from openpyxl import load_workbook

from app.mcp.excel_client import ExcelMCPClient, ExcelMCPError


def test_sync_and_read_through_mcp(tmp_path):
    path = tmp_path / "plans.xlsx"

    async def scenario():
        async with ExcelMCPClient(path).session() as excel:
            assert {"get_flights", "get_hotels", "get_activities", "get_all_travel_options",
                    "sync_travel_options", "get_cost_summary"} <= set(await excel.list_tools())
            await excel.sync_travel_options(
                "s1",
                flights=[{"direction": "outbound", "airline": "A", "price": 100, "currency": "USD"}],
                hotels=[{"name": "H", "total_price": 500, "currency": "USD"}],
                activities=[{"name": "Museum", "price": 20, "currency": "USD", "day": 1}],
            )
            await excel.sync_travel_options("s2", flights=[], hotels=[{"name": "Other"}], activities=[])
            # Re-syncing a session replaces its rows instead of duplicating them.
            await excel.sync_travel_options(
                "s1",
                flights=[{"direction": "outbound", "airline": "B", "price": 90, "currency": "USD"}],
                hotels=[{"name": "H", "total_price": 500, "currency": "USD"}],
                activities=[],
            )
            return (await excel.get_all_travel_options("s1"), await excel.get_hotels("s2"),
                    await excel.get_cost_summary("s1"))

    s1, s2_hotels, summary = asyncio.run(scenario())
    assert [f["airline"] for f in s1["flights"]] == ["B"]
    assert s1["hotels"][0]["name"] == "H"
    assert s1["activities"] == []
    assert [h["name"] for h in s2_hotels] == ["Other"]
    assert summary["flights"]["outbound:USD"]["min"] == 90

    wb = load_workbook(path)
    assert {"Flights", "Hotels", "Activities"} <= set(wb.sheetnames)


def test_connection_failure_raises_excel_mcp_error(tmp_path):
    async def scenario():
        async with ExcelMCPClient(tmp_path / "x.xlsx", python_executable="no-such-python-exe").session():
            pass

    with pytest.raises(ExcelMCPError):
        asyncio.run(scenario())

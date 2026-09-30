"""End-to-end graph tests. They use the real Excel MCP server (stdio subprocess) unless noted."""
import asyncio
import time
from contextlib import asynccontextmanager

from openpyxl import load_workbook

from app.agents import activity_agent, flight_agent, hotel_agent
from app.blackboard import BlackboardKeys
from app.mcp.excel_client import ExcelMCPClient
from app.tools.common import ToolError


def test_full_workflow_produces_plan(run_trip, paris_request, cfg):
    result = run_trip(paris_request)
    state, bb = result.state, result.blackboard

    assert state["request_valid"]
    assert (state["flight_status"], state["hotel_status"], state["activity_status"]) == ("ok", "ok", "ok")
    assert state["excel_sync_status"] == "ok"
    assert state["errors"] == []

    # Specialists wrote to the Blackboard, and the planner seeded the constraints.
    for key in (BlackboardKeys.TRIP_CONSTRAINTS, *BlackboardKeys.SPECIALIST_KEYS):
        assert bb.has(key)
    assert bb.entry(BlackboardKeys.FLIGHTS).written_by == "flight_agent"

    # The specialist data reached Excel, under this session's id.
    wb = load_workbook(cfg.EXCEL_WORKBOOK_PATH)
    hotel_rows = [r for r in wb["Hotels"].iter_rows(values_only=True) if r[0] == bb.session_id]
    assert len(hotel_rows) == len(bb.read(BlackboardKeys.HOTELS)["options"])

    budget = state["budget_result"]
    assert budget["status"] == "ok"
    assert budget["selected"]["estimated_total"] <= 1500
    assert budget["selected"]["flight_cost"] == (
        budget["selected"]["outbound_flight"]["price"] + budget["selected"]["return_flight"]["price"]
    )

    md = result.plan.markdown
    for heading in ("# Travel Plan", "## Flight", "## Hotel", "## Activities", "**Day 1", "## Budget", "## Notes"):
        assert heading in md
    assert "demo catalog" in md


def test_specialists_run_in_parallel(run_trip, paris_request, monkeypatch):
    spans = {}

    def slow(module, name):
        original = getattr(module, name)

        async def wrapper(constraints, cfg=None):
            start = time.perf_counter()
            await asyncio.sleep(0.4)
            result = await original(constraints, cfg)
            spans[name] = (start, time.perf_counter())
            return result

        monkeypatch.setattr(module, name, wrapper)

    slow(flight_agent, "search_flights")
    slow(hotel_agent, "search_hotels")
    slow(activity_agent, "search_activities")

    run_trip(paris_request)

    assert len(spans) == 3
    latest_start = max(s for s, _ in spans.values())
    earliest_end = min(e for _, e in spans.values())
    assert latest_start < earliest_end, f"specialists did not overlap: {spans}"


def test_flight_failure_does_not_stop_workflow(run_trip, paris_request, monkeypatch):
    async def broken(constraints, cfg=None):
        raise ToolError("flight provider is down")

    monkeypatch.setattr(flight_agent, "search_flights", broken)
    result = run_trip(paris_request)
    state = result.state

    assert state["flight_status"] == "failed"
    assert state["hotel_status"] == "ok" and state["activity_status"] == "ok"
    assert result.blackboard.read(BlackboardKeys.FLIGHTS)["status"] == "failed"

    budget = state["budget_result"]
    assert budget["status"] == "partial"
    assert "flights" in budget["unavailable_components"]
    assert budget["selected"]["flight_cost"] is None
    assert budget["selected"]["hotel_cost"] is not None

    md = result.plan.markdown
    assert "Flight information is unavailable" in md
    assert "flight provider is down" in md


def test_mcp_failure_marks_budget_unavailable(run_trip, paris_request, cfg):
    broken_client = ExcelMCPClient(cfg.EXCEL_WORKBOOK_PATH, python_executable="no-such-python-exe")
    result = run_trip(paris_request, excel_client=broken_client)
    state = result.state

    assert state["excel_sync_status"] == "failed"
    assert state["budget_result"]["status"] == "unavailable"
    # Specialist results are still reported from the Blackboard.
    assert "## Hotel" in result.plan.markdown and "Montmartre Budget Inn" in result.plan.markdown
    assert "Budget estimate unavailable" in result.plan.markdown


def test_invalid_request_skips_specialists(run_trip):
    result = run_trip({"travel_date": "2020-01-01", "duration_days": 0})
    state = result.state

    assert state["request_valid"] is False
    assert "flight_status" not in state and "budget_result" not in state
    assert not result.blackboard.has(BlackboardKeys.TRIP_CONSTRAINTS)
    md = result.plan.markdown
    assert "Destination is missing." in md
    assert "Travel date is invalid" in md
    assert "Duration is invalid" in md


def test_free_text_without_llm_is_rejected_cleanly(run_trip):
    result = run_trip("5 days in Paris next month")
    assert result.state["request_valid"] is False
    assert "LLM_API_KEY" in result.plan.markdown


def test_missing_origin_skips_flights(run_trip, paris_request):
    paris_request.pop("origin")
    result = run_trip(paris_request)

    assert result.state["flight_status"] == "skipped"
    assert "flights" in result.state["budget_result"]["unavailable_components"]
    assert "No origin was provided" in result.plan.markdown


def test_unknown_destination_reports_no_results(run_trip, paris_request):
    paris_request["destination"] = "Atlantis"
    result = run_trip(paris_request)

    assert {result.state[k] for k in ("flight_status", "hotel_status", "activity_status")} == {"empty"}
    assert result.state["budget_result"]["status"] == "unavailable"
    assert "No hotels found in Atlantis" in result.plan.markdown


class FakeExcelSession:
    """Stands in for the MCP server. Returns a hotel that never appears on the Blackboard."""

    def __init__(self):
        self.synced = None

    async def sync_travel_options(self, session_id, flights, hotels, activities):
        self.synced = {"flights": flights, "hotels": hotels, "activities": activities}
        return {"rows_written": {}}

    async def get_flights(self, session_id):
        return self.synced["flights"]

    async def get_hotels(self, session_id):
        return [{"name": "Only-In-Excel Hotel", "location": "X", "rating": 5, "price_per_night": 1,
                 "total_price": 5, "currency": "USD"}]

    async def get_activities(self, session_id):
        return self.synced["activities"]


class FakeExcelClient:
    def __init__(self):
        self.fake = FakeExcelSession()

    @asynccontextmanager
    async def session(self):
        yield self.fake


def test_budget_agent_reads_from_excel_mcp_not_blackboard(run_trip, paris_request):
    result = run_trip(paris_request, excel_client=FakeExcelClient())
    assert result.state["budget_result"]["selected"]["hotel"]["name"] == "Only-In-Excel Hotel"

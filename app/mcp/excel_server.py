"""Excel MCP server: exposes the travel-plans workbook as MCP tools.

The workbook has one sheet per specialist (Flights, Hotels, Activities). Every
row carries a `session_id` column, so several planning sessions can share one
file. All openpyxl details live here. Agents only see the MCP tools.

Run standalone (stdio transport):
    python -m app.mcp.excel_server
"""
import os
import tempfile
import time
from collections import defaultdict
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from mcp.server.fastmcp import FastMCP
from openpyxl import Workbook, load_workbook

DEFAULT_WORKBOOK = Path(__file__).resolve().parents[2] / "data" / "travel_plans.xlsx"
WORKBOOK_PATH = Path(os.environ.get("EXCEL_WORKBOOK_PATH") or DEFAULT_WORKBOOK)

SHEETS: dict[str, list[str]] = {
    "Flights": [
        "session_id", "direction", "airline", "flight_number", "departure", "arrival",
        "departure_time", "arrival_time", "duration", "price", "currency", "source", "synced_at",
    ],
    "Hotels": [
        "session_id", "name", "location", "rating", "price_per_night", "total_price",
        "currency", "source", "synced_at",
    ],
    "Activities": [
        "session_id", "day", "name", "category", "description", "duration_hours",
        "price", "currency", "source", "synced_at",
    ],
}
SHEET_FOR_KIND = {"flights": "Flights", "hotels": "Hotels", "activities": "Activities"}
PRICE_FIELD = {"flights": "price", "hotels": "total_price", "activities": "price"}

LOCK_TIMEOUT_SECONDS = 15
STALE_LOCK_SECONDS = 60

mcp = FastMCP("excel-travel-plans", log_level="WARNING")


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------
@contextmanager
def _workbook_lock(path: Path):
    """Cross-process lock so concurrent sessions never corrupt the workbook."""
    lock_path = path.with_suffix(path.suffix + ".lock")
    deadline = time.monotonic() + LOCK_TIMEOUT_SECONDS
    while True:
        try:
            fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            break
        except FileExistsError:
            try:
                if time.time() - lock_path.stat().st_mtime > STALE_LOCK_SECONDS:
                    lock_path.unlink(missing_ok=True)
                    continue
            except FileNotFoundError:
                continue
            if time.monotonic() > deadline:
                raise TimeoutError(f"Timed out waiting for workbook lock {lock_path}")
            time.sleep(0.05)
    try:
        yield
    finally:
        os.close(fd)
        lock_path.unlink(missing_ok=True)


def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns in SHEETS.items():
        wb.create_sheet(name).append(columns)
    return wb


def _open_workbook(path: Path) -> Workbook:
    if not path.exists():
        return _new_workbook()
    wb = load_workbook(path)
    for name, columns in SHEETS.items():
        if name not in wb.sheetnames:
            wb.create_sheet(name).append(columns)
    return wb


def _save_atomic(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _read_rows(kind: str, session_id: str) -> list[dict[str, Any]]:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook {WORKBOOK_PATH} does not exist yet; nothing has been synced")
    wb = load_workbook(WORKBOOK_PATH, read_only=True)
    try:
        sheet_name = SHEET_FOR_KIND[kind]
        if sheet_name not in wb.sheetnames:
            return []
        rows = wb[sheet_name].iter_rows(values_only=True)
        header = list(next(rows, []))
        records = [dict(zip(header, row)) for row in rows]
    finally:
        wb.close()
    return [
        {k: v for k, v in record.items() if k not in ("session_id", "synced_at")}
        for record in records
        if record.get("session_id") == session_id
    ]


def _replace_session_rows(wb: Workbook, sheet_name: str, session_id: str, rows: list[dict[str, Any]]) -> int:
    ws = wb[sheet_name]
    columns = SHEETS[sheet_name]
    for idx in range(ws.max_row, 1, -1):
        if ws.cell(row=idx, column=1).value == session_id:
            ws.delete_rows(idx)
    synced_at = datetime.now(timezone.utc).isoformat()
    for row in rows:
        record = {**row, "session_id": session_id, "synced_at": synced_at}
        ws.append([record.get(col) for col in columns])
    return len(rows)


# ---------------------------------------------------------------------------
# MCP tools
# ---------------------------------------------------------------------------
@mcp.tool()
def sync_travel_options(
    session_id: str,
    flights: list[dict[str, Any]],
    hotels: list[dict[str, Any]],
    activities: list[dict[str, Any]],
) -> dict[str, Any]:
    """Write (replace) this session's flight, hotel and activity options into the workbook."""
    with _workbook_lock(WORKBOOK_PATH):
        wb = _open_workbook(WORKBOOK_PATH)
        written = {
            "flights": _replace_session_rows(wb, "Flights", session_id, flights),
            "hotels": _replace_session_rows(wb, "Hotels", session_id, hotels),
            "activities": _replace_session_rows(wb, "Activities", session_id, activities),
        }
        _save_atomic(wb, WORKBOOK_PATH)
    return {"session_id": session_id, "workbook": str(WORKBOOK_PATH), "rows_written": written}


@mcp.tool()
def get_flights(session_id: str) -> dict[str, Any]:
    """Return the flight options (outbound and return legs) stored for a session."""
    return {"session_id": session_id, "flights": _read_rows("flights", session_id)}


@mcp.tool()
def get_hotels(session_id: str) -> dict[str, Any]:
    """Return the hotel options stored for a session."""
    return {"session_id": session_id, "hotels": _read_rows("hotels", session_id)}


@mcp.tool()
def get_activities(session_id: str) -> dict[str, Any]:
    """Return the activity options stored for a session."""
    return {"session_id": session_id, "activities": _read_rows("activities", session_id)}


@mcp.tool()
def get_all_travel_options(session_id: str) -> dict[str, Any]:
    """Return flights, hotels and activities stored for a session in one call."""
    return {
        "session_id": session_id,
        "flights": _read_rows("flights", session_id),
        "hotels": _read_rows("hotels", session_id),
        "activities": _read_rows("activities", session_id),
    }


@mcp.tool()
def get_cost_summary(session_id: str) -> dict[str, Any]:
    """Min / max / count of prices per category (flights split by direction), grouped by currency."""
    summary: dict[str, Any] = {}
    for kind in ("flights", "hotels", "activities"):
        groups: dict[str, list[float]] = defaultdict(list)
        for row in _read_rows(kind, session_id):
            price = row.get(PRICE_FIELD[kind])
            if price is None:
                continue
            group = f"{row['direction']}:{row.get('currency')}" if kind == "flights" else str(row.get("currency"))
            groups[group].append(float(price))
        summary[kind] = {
            group: {"count": len(prices), "min": min(prices), "max": max(prices)}
            for group, prices in groups.items()
        }
    return {"session_id": session_id, "summary": summary}


if __name__ == "__main__":
    mcp.run(transport="stdio")
